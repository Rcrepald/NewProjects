"""Exporta tudo para uma planilha Excel (backup legível)."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from . import servico as s


def exportar(conn, hoje: date, destino: Path) -> Path:
    wb = Workbook()

    def aba(titulo, cabecalho, linhas, primeira=False):
        ws = wb.active if primeira else wb.create_sheet()
        ws.title = titulo
        ws.append(cabecalho)
        for c in ws[1]:
            c.font = Font(bold=True)
        for linha in linhas:
            ws.append(linha)
        for col in ws.columns:
            largura = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(largura + 2, 45)

    alunos = conn.execute("SELECT * FROM alunos ORDER BY nome").fetchall()
    aba("Alunos", ["ID", "Nome", "Telefone", "Plano", "Mensalidade (R$)", "Dia venc.", "Cobra desde",
                   "Ativo", "Observações"],
        [[a["id"], a["nome"], a["telefone"], a["plano"], a["valor"] / 100, a["dia_vencimento"],
          a["inicio_cobranca"], "sim" if a["ativo"] else "não", a["observacoes"]] for a in alunos],
        primeira=True)

    ms = s.mensalidades_com_saldo(conn)
    aba("Mensalidades", ["Aluno", "Mês", "Vencimento", "Valor (R$)", "Pago (R$)", "Falta (R$)", "Situação"],
        [[m["nome"], m["competencia"], m["vencimento"], m["valor"] / 100, m["pago"] / 100, m["saldo"] / 100,
          s.situacao(m, m["pago"], hoje)] for m in ms])

    pags = conn.execute(
        "SELECT p.*, a.nome, m.competencia FROM pagamentos p JOIN alunos a ON a.id = p.aluno_id"
        " LEFT JOIN mensalidades m ON m.id = p.mensalidade_id ORDER BY p.data, p.id").fetchall()
    aba("Pagamentos", ["Data", "Aluno", "Valor (R$)", "Forma", "Referente a", "Observação"],
        [[p["data"], p["nome"], p["valor"] / 100, p["forma"], p["competencia"], p["observacao"]] for p in pags])

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
